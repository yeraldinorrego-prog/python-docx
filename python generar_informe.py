from docx import Document
from docx.shared import Pt

# Crear documento Word
doc = Document()
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)

# Portada
doc.add_paragraph('Radar de Oportunidades de Mercado', 'Title').alignment = 1
doc.add_paragraph('Producto: Concentrado Nutritivo para Animales de Granja').alignment = 1
doc.add_paragraph('Materia: [Nombre de la asignatura]').alignment = 1
doc.add_paragraph('Integrantes: [Nombre del equipo o integrantes]').alignment = 1
doc.add_paragraph('Fecha: [Fecha de entrega]').alignment = 1
doc.add_paragraph('Institución: [Nombre de la institución]').alignment = 1
doc.add_page_break()

# Tabla de contenido
doc.add_heading('Tabla de Contenido', level=1)
for i, title in enumerate([
    "Introducción", "Descripción del Producto", "Análisis del Entorno (PESTEL)",
    "Segmento de Mercado", "Diseño del Marketing Mix", "Validación de la Propuesta",
    "Conclusiones", "Referencias"
], 1):
    doc.add_paragraph(f"{i}. {title}")
doc.add_page_break()

# (Agrega aquí los mismos contenidos de las secciones que ya te redacté)
# Para mantenerlo breve, copia el contenido de cada sección desde la respuesta anterior.

# Guardar archivo
doc.save("Radar_de_Oportunidades_Concentrado_Animales_APA.docx")
